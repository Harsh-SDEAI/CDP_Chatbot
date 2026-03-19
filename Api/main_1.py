from dotenv import load_dotenv
load_dotenv()
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import uuid
import httpx
from bs4 import BeautifulSoup
from pathlib import Path
import json
from datetime import datetime
import re
import shutil
import numpy as np
from typing import Any, Dict, List, Optional
from openai import OpenAI

# ── pip install docling faiss-cpu openai
from docling.document_converter import DocumentConverter, PdfFormatOption, ImageFormatOption
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions
from docling.datamodel.pipeline_options import EasyOcrOptions
import faiss

app = FastAPI(title="Admin Content Manager API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STORAGE_DIR = Path("storage")
STORAGE_DIR.mkdir(exist_ok=True)

IMAGES_DIR = STORAGE_DIR / "images"
IMAGES_DIR.mkdir(exist_ok=True)

FAISS_DIR = STORAGE_DIR / "faiss"
FAISS_DIR.mkdir(exist_ok=True)

app.mount("/images", StaticFiles(directory=str(IMAGES_DIR)), name="images")

# ── OpenAI client
ai = OpenAI()

# ── FAISS index (in-memory, persisted to disk)
FAISS_INDEX_PATH = FAISS_DIR / "index.faiss"
FAISS_META_PATH  = FAISS_DIR / "index_meta.json"
EMBEDDING_DIM    = 1536  # text-embedding-3-small

# Global index + metadata list
_faiss_index: Optional[faiss.IndexFlatL2] = None
_faiss_meta: List[Dict[str, Any]] = []   # [{file_id, chunk_index, text}, ...]


# ─── Models ───────────────────────────────────────────────────────────────────

class ScrapeRequest(BaseModel):
    url: str


class RAGQueryRequest(BaseModel):
    query: str
    top_k: int = 5


# ─── FAISS Helpers ────────────────────────────────────────────────────────────

def _load_faiss():
    global _faiss_index, _faiss_meta
    if FAISS_INDEX_PATH.exists() and FAISS_META_PATH.exists():
        _faiss_index = faiss.read_index(str(FAISS_INDEX_PATH))
        _faiss_meta = json.loads(FAISS_META_PATH.read_text(encoding="utf-8"))
        print(f"[FAISS] Loaded index with {_faiss_index.ntotal} vectors.")
    else:
        _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
        _faiss_meta = []
        print("[FAISS] Created new empty index.")


def _save_faiss():
    faiss.write_index(_faiss_index, str(FAISS_INDEX_PATH))
    FAISS_META_PATH.write_text(json.dumps(_faiss_meta, ensure_ascii=False), encoding="utf-8")


def _get_embedding(text: str) -> np.ndarray:
    response = ai.embeddings.create(
        model="text-embedding-3-small",
        input=text[:8000],  # stay within token limit
    )
    return np.array(response.data[0].embedding, dtype=np.float32)


def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """Split text into overlapping chunks."""
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
        i += chunk_size - overlap
    return chunks


def index_document(file_id: str, content: str):
    """Chunk a document, embed each chunk, add to FAISS index."""
    global _faiss_index, _faiss_meta

    # Remove any existing chunks for this file_id first
    remove_document_from_index(file_id)

    chunks = _chunk_text(content)
    vectors = []
    metas = []

    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        try:
            vec = _get_embedding(chunk)
            vectors.append(vec)
            metas.append({
                "file_id": file_id,
                "chunk_index": i,
                "text": chunk,
            })
        except Exception as e:
            print(f"[FAISS] Embedding failed for chunk {i} of {file_id}: {e}")

    if vectors:
        matrix = np.stack(vectors)
        _faiss_index.add(matrix)
        _faiss_meta.extend(metas)
        _save_faiss()
        print(f"[FAISS] Indexed {len(vectors)} chunks for {file_id}.")


def remove_document_from_index(file_id: str):
    """Remove all chunks for a file_id by rebuilding the index without them."""
    global _faiss_index, _faiss_meta

    remaining = [m for m in _faiss_meta if m["file_id"] != file_id]
    if len(remaining) == len(_faiss_meta):
        return  # nothing to remove

    # Rebuild index from remaining chunks
    _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
    _faiss_meta = []

    if remaining:
        vectors = []
        for m in remaining:
            try:
                vec = _get_embedding(m["text"])
                vectors.append(vec)
                _faiss_meta.append(m)
            except Exception:
                pass
        if vectors:
            _faiss_index.add(np.stack(vectors))

    _save_faiss()
    print(f"[FAISS] Removed chunks for {file_id}, index now has {_faiss_index.ntotal} vectors.")


def search_index(query: str, top_k: int = 5) -> List[Dict[str, Any]]:
    """Search FAISS index and return top_k matching chunks with metadata."""
    if _faiss_index is None or _faiss_index.ntotal == 0:
        return []

    query_vec = _get_embedding(query)
    query_matrix = np.array([query_vec], dtype=np.float32)

    k = min(top_k, _faiss_index.ntotal)
    distances, indices = _faiss_index.search(query_matrix, k)

    results = []
    for dist, idx in zip(distances[0], indices[0]):
        if idx == -1:
            continue
        meta = _faiss_meta[idx].copy()
        meta["score"] = float(dist)
        results.append(meta)

    return results


# ─── Storage Helpers ──────────────────────────────────────────────────────────

def save_file(file_id: str, content: str) -> Path:
    path = STORAGE_DIR / f"{file_id}.md"
    path.write_text(content, encoding="utf-8")
    return path


def load_file(file_id: str) -> str:
    path = STORAGE_DIR / f"{file_id}.md"
    if not path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return path.read_text(encoding="utf-8")


def save_meta(file_id: str, meta: dict):
    path = STORAGE_DIR / f"{file_id}.meta.json"
    path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")


def load_meta(file_id: str) -> dict:
    path = STORAGE_DIR / f"{file_id}.meta.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(file_id: str, data: dict) -> Path:
    path = STORAGE_DIR / f"{file_id}.json"
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def _make_file_id(filename: str) -> str:
    stem = Path(filename).stem
    safe = re.sub(r"[^A-Za-z0-9_\-]", "_", stem)
    safe = re.sub(r"_+", "_", safe).strip("_")
    candidate = safe
    counter = 2
    while (STORAGE_DIR / f"{candidate}.md").exists():
        candidate = f"{safe}_{counter}"
        counter += 1
    return candidate


def _make_scrape_filename(url: str) -> str:
    """Generate a human-readable filename from a URL.

    Examples:
        https://www.cooperstowndreamspark.com/testimonials/  -> cooperstowndreamspark_com_testimonials
        https://example.com                                  -> example_com
        https://xyz.com/abc/def                              -> xyz_com_abc_def
        https://xyz.com/abc/def.html                         -> xyz_com_abc_def
    """
    from urllib.parse import urlparse
    parsed = urlparse(url)
    domain = parsed.netloc.replace("www.", "")
    path   = parsed.path.strip("/")

    if path:
        segments = path.split("/")
        segments = [Path(s).stem if "." in s else s for s in segments]
        name = domain + "/" + "/".join(segments)
    else:
        name = domain

    safe = re.sub(r"[^A-Za-z0-9]", "_", name)
    safe = re.sub(r"_+", "_", safe).strip("_")

    candidate = safe
    counter = 2
    while (STORAGE_DIR / f"{candidate}.md").exists():
        candidate = f"{safe}_{counter}"
        counter += 1
    return candidate


def html_to_markdown(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    lines = []
    for elem in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote"]):
        tag = elem.name
        text = elem.get_text(separator=" ", strip=True)
        if not text:
            continue
        if tag == "h1":           lines.append(f"# {text}")
        elif tag == "h2":         lines.append(f"## {text}")
        elif tag == "h3":         lines.append(f"### {text}")
        elif tag == "h4":         lines.append(f"#### {text}")
        elif tag == "li":         lines.append(f"- {text}")
        elif tag == "pre":        lines.append(f"```\n{text}\n```")
        elif tag == "blockquote": lines.append(f"> {text}")
        else:                     lines.append(text)
        lines.append("")
    return "\n".join(lines).strip()


# ─── Docling Helpers ──────────────────────────────────────────────────────────

BULLET_CHARS   = ["•", "·", "●", "▪", "", "o", "◦", "-"]
CHECKBOX_CHARS = ["□", "☐", "✓", "✔", "✗", "✘", ""]


def _get_by_ref(doc: Dict[str, Any], ref_str: str) -> Any:
    if not ref_str.startswith("#/"):
        raise ValueError(f"Bad ref: {ref_str}")
    parts = ref_str[2:].split("/")
    obj: Any = doc
    for p in parts:
        obj = obj[int(p)] if isinstance(obj, list) else obj[p]
    return obj


def _extract_pages(text_obj: Dict[str, Any]) -> set:
    pages = set()
    for prov in text_obj.get("prov", []):
        if prov.get("page_no") is not None:
            pages.add(int(prov["page_no"]))
    return pages


def _clean_bullet(s: str) -> str:
    s = (s or "").lstrip()
    if not s:
        return s
    for ch in BULLET_CHARS + CHECKBOX_CHARS:
        if s.startswith(ch):
            return s[len(ch):].lstrip()
    return re.sub(r"^[^\w\d]+", "", s).lstrip()


def _format_list_item(obj: Dict[str, Any], indent: int) -> str:
    text   = (obj.get("text") or "").strip()
    orig   = (obj.get("orig") or "").strip()
    marker = (obj.get("marker") or "").strip()
    pad    = "    " * indent
    if obj.get("enumerated") and marker:
        return pad + f"{marker} {text}"
    base = _clean_bullet(orig) or text
    return pad + f"- {base}"


def _format_checkbox(obj: Dict[str, Any], indent: int) -> str:
    orig  = (obj.get("orig") or "").strip()
    text  = (obj.get("text") or "").strip()
    pad   = "    " * indent
    base  = _clean_bullet(orig) or _clean_bullet(text) or text
    label = obj.get("label") or ""
    box   = "[x]" if "selected" in label else "[ ]"
    return f"{pad}{box} {base}"


def _docling_doc_to_markdown(doc_dict: Dict[str, Any]) -> str:
    lines: List[str] = []

    def _render_table(table_obj):
        grid = table_obj.get("grid") or table_obj.get("data", {}).get("grid", [])
        if not grid:
            return []

        # Check if this is a menu/document_index style table
        # by detecting price patterns ($xx) in cells
        all_texts = [
            (cell.get("text") or "").strip()
            for row in grid for cell in row
        ]
        is_menu = any(re.match(r'^\$\d+', t) for t in all_texts)

        if is_menu:
            tlines = []
            seen_texts = set()
            for row in grid:
                # Deduplicate spanned cells (same text repeated due to col_span)
                unique_cells = []
                for cell in row:
                    t = (cell.get("text") or "").strip()
                    if t and t not in seen_texts:
                        seen_texts.add(t)
                        unique_cells.append(t)

                if not unique_cells:
                    continue

                # Separate price from the rest
                price = next((t for t in unique_cells if re.match(r'^\$\d+', t)), None)
                content_cells = [t for t in unique_cells if not re.match(r'^\$\d+', t)]
                content = " | ".join(content_cells) if content_cells else ""

                if content and price:
                    tlines.append(f"- {content} {price}")
                elif content:
                    tlines.append(f"- {content}")

            return tlines

        # Default: render as markdown table
        tlines = []
        for row_idx, row in enumerate(grid):
            cells = [str((cell.get("text") or "").strip()) for cell in row]
            tlines.append("| " + " | ".join(cells) + " |")
            if row_idx == 0:
                tlines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        return tlines

    def walk_ref(ref: str, indent: int = 0):
        kind = ref.split("/")[1]
        try:
            obj = _get_by_ref(doc_dict, ref)
        except Exception:
            return

        if kind == "texts":
            if obj.get("content_layer") != "body":
                return
            label = obj.get("label", "")
            text  = (obj.get("text") or "").strip()
            if not text:
                return
            if label == "section_header":
                lines.append(f"\n{'#' * obj.get('level', 2)} {text}\n")
            elif label == "list_item":
                lines.append(_format_list_item(obj, indent))
            elif label and label.startswith("checkbox"):
                lines.append(_format_checkbox(obj, indent))
            else:
                lines.append(text)

        elif kind == "groups":
            child_indent = indent + 1 if obj.get("label") == "list" else indent
            for child in obj.get("children", []):
                walk_ref(child["$ref"], child_indent)

        elif kind == "tables":
            tlines = _render_table(obj)
            if tlines:
                lines.append("")
                lines.extend(tlines)
                lines.append("")

        elif kind == "pictures":
            captions, ocr_texts = [], []
            for ch in obj.get("children", []):
                try:
                    t     = _get_by_ref(doc_dict, ch["$ref"])
                    text  = (t.get("text") or "").strip()
                    label = (t.get("label") or "").strip()
                    if not text:
                        continue
                    (captions if label in ("caption", "footnote") else ocr_texts).append(text)
                except Exception:
                    pass
            parts = []
            if captions:  parts.append("Caption: " + " ".join(captions))
            if ocr_texts: parts.append("Image text: " + " ".join(ocr_texts))
            if parts:
                lines.append(f"*[Image — {' | '.join(parts)}]*")

    for child in doc_dict.get("body", {}).get("children", []):
        walk_ref(child["$ref"], indent=0)
        lines.append("")

    return "\n".join(lines).strip()


def _docling_doc_to_structured_json(doc_dict: Dict[str, Any], file_id: str, filename: str) -> Dict[str, Any]:
    blocks: List[Dict[str, Any]] = []

    def walk_ref(ref: str, current_section: Optional[str] = None):
        kind = ref.split("/")[1]
        try:
            obj = _get_by_ref(doc_dict, ref)
        except Exception:
            return
        pages = sorted(_extract_pages(obj))

        if kind == "texts":
            if obj.get("content_layer") != "body":
                return
            text = (obj.get("text") or "").strip()
            if not text:
                return
            blocks.append({"type": obj.get("label", "paragraph"), "text": text, "section": current_section, "pages": pages})

        elif kind == "groups":
            items, g_pages = [], set()
            for ch in obj.get("children", []):
                try:
                    t = _get_by_ref(doc_dict, ch["$ref"])
                    text = (t.get("text") or "").strip()
                    if text:
                        items.append(text)
                        g_pages.update(_extract_pages(t))
                except Exception:
                    pass
            if items:
                blocks.append({"type": obj.get("label", "group"), "items": items, "section": current_section, "pages": sorted(g_pages)})

        elif kind == "tables":
            grid = obj.get("grid", [])
            rows = [[str((c.get("text") or "").strip()) for c in row] for row in grid]
            if rows:
                blocks.append({"type": "table", "headers": rows[0], "rows": rows[1:], "section": current_section, "pages": pages})

        elif kind == "pictures":
            captions, ocr_texts, p_pages = [], [], set(pages)
            for ch in obj.get("children", []):
                try:
                    t     = _get_by_ref(doc_dict, ch["$ref"])
                    text  = (t.get("text") or "").strip()
                    label = (t.get("label") or "").strip()
                    if not text: continue
                    p_pages.update(_extract_pages(t))
                    (captions if label in ("caption", "footnote") else ocr_texts).append(text)
                except Exception:
                    pass
            if captions or ocr_texts:
                blocks.append({"type": "picture", "captions": captions, "ocr_text": ocr_texts, "section": current_section, "pages": sorted(p_pages)})

    current_section = None
    for child in doc_dict.get("body", {}).get("children", []):
        ref = child["$ref"]
        if ref.split("/")[1] == "texts":
            try:
                obj = _get_by_ref(doc_dict, ref)
                if obj.get("label") == "section_header":
                    current_section = (obj.get("text") or "").strip()
            except Exception:
                pass
        walk_ref(ref, current_section)

    origin = doc_dict.get("origin", {})
    return {
        "file_id": file_id,
        "filename": filename,
        "page_count": len(doc_dict.get("pages", {})),
        "extracted_at": datetime.utcnow().isoformat() + "Z",
        "source": {"mime_type": origin.get("mimetype"), "original_filename": origin.get("filename")},
        "blocks": blocks,
    }


def image_to_docling(img_path: Path, file_id: str, filename: str):
    print(f"[Docling] Starting image OCR extraction for {file_id}...")
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = True
    pipeline_options.do_table_structure = True
    pipeline_options.images_scale = 2.0
    pipeline_options.ocr_options = EasyOcrOptions(lang=["en"])


    converter = DocumentConverter(
        format_options={InputFormat.IMAGE: ImageFormatOption(pipeline_options=pipeline_options)}
    )
    result   = converter.convert(str(img_path))
    doc_dict = result.document.export_to_dict()
    print(f"[Docling] Image extraction complete for {file_id}.")

    markdown       = _docling_doc_to_markdown(doc_dict)
    structured_json = _docling_doc_to_structured_json(doc_dict, file_id, filename)
    structured_json["docling_raw"] = doc_dict
    return markdown, structured_json


def pdf_to_docling(pdf_path: Path, file_id: str, filename: str):
    print(f"[Docling] Starting full OCR extraction for {file_id}...")
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = True
    pipeline_options.do_table_structure = True
    pipeline_options.images_scale = 2.0
    pipeline_options.generate_picture_images = True

    converter = DocumentConverter(
        format_options={"pdf": PdfFormatOption(pipeline_options=pipeline_options)}
    )
    result   = converter.convert(str(pdf_path))
    doc_dict = result.document.export_to_dict()
    print(f"[Docling] Extraction complete for {file_id}.")

    markdown        = _docling_doc_to_markdown(doc_dict)
    structured_json = _docling_doc_to_structured_json(doc_dict, file_id, filename)
    structured_json["docling_raw"] = doc_dict
    return markdown, structured_json


def docx_to_markdown(docx_path: Path, file_id: str, filename: str):
    """Convert DOCX to markdown using python-docx."""
    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    print(f"[DOCX] Starting extraction for {file_id}...")
    doc = DocxDocument(str(docx_path))
    lines: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        style_name = (para.style.name or "").lower()

        # ── Headings ──
        if style_name.startswith("heading"):
            try:
                level = int(style_name.replace("heading", "").strip())
            except ValueError:
                level = 1
            lines.append(f"{'#' * level} {text}")

        # ── Bullet / numbered lists ──
        elif style_name.startswith("list"):
            # Check if it's a numbered list
            if "number" in style_name or "ordered" in style_name:
                lines.append(f"1. {text}")
            else:
                lines.append(f"* {text}")

        # ── Normal paragraph (apply bold/italic from runs) ──
        else:
            rich_parts: List[str] = []
            for run in para.runs:
                t = run.text
                if not t:
                    continue
                if run.bold and run.italic:
                    t = f"***{t}***"
                elif run.bold:
                    t = f"**{t}**"
                elif run.italic:
                    t = f"*{t}*"
                rich_parts.append(t)
            lines.append("".join(rich_parts) if rich_parts else text)

    # ── Tables ──
    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        lines.append("")

    markdown = "\n".join(lines).strip()
    print(f"[DOCX] Extraction complete for {file_id} ({len(markdown)} chars).")
    return markdown


# ─── Startup ──────────────────────────────────────────────────────────────────

@app.on_event("startup")
async def startup_event():
    _load_faiss()


# ─── Routes ───────────────────────────────────────────────────────────────────

@app.get("/")
def root():
    return {"message": "Admin Content Manager API is running"}


@app.post("/scrape")
async def scrape_url(body: ScrapeRequest):
    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20) as client:
            response = await client.get(
                body.url,
                headers={"User-Agent": "Mozilla/5.0 (compatible; AdminContentBot/1.0)"}
            )
            response.raise_for_status()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to fetch URL: {str(e)}")

    markdown = html_to_markdown(response.text)
    file_id  = _make_scrape_filename(body.url)

    save_file(file_id, markdown)
    save_meta(file_id, {
        "type": "scraped",
        "url": body.url,
        "filename": f"{file_id}.md",
        "created_at": datetime.utcnow().isoformat() + "Z"
    })

    # Index for RAG
    index_document(file_id, markdown)

    return {"file_id": file_id, "content": markdown, "filename": f"{file_id}.md", "url": body.url, "source_url": body.url}


@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    filename = file.filename
    ext      = Path(filename).suffix.lower()

    IMAGE_EXTS   = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}
    DOCX_EXTS    = {".docx", ".doc"}
    ALLOWED_EXTS = {".pdf", ".txt", ".md"} | IMAGE_EXTS | DOCX_EXTS

    if ext not in ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

    raw_content = await file.read()
    file_id     = _make_file_id(filename)

    if ext == ".pdf":
        tmp_pdf = STORAGE_DIR / f"{file_id}_tmp.pdf"
        tmp_pdf.write_bytes(raw_content)
        try:
            markdown, structured_json = pdf_to_docling(tmp_pdf, file_id, filename)
        finally:
            tmp_pdf.unlink(missing_ok=True)

        save_file(file_id, markdown)
        json_path = save_json(file_id, structured_json)
        content   = markdown
        extra     = {"json_path": str(json_path), "page_count": structured_json.get("page_count", 0), "block_count": len(structured_json.get("blocks", []))}

    elif ext in IMAGE_EXTS:
        tmp_img = STORAGE_DIR / f"{file_id}_tmp{ext}"
        tmp_img.write_bytes(raw_content)
        try:
            markdown, structured_json = image_to_docling(tmp_img, file_id, filename)
        finally:
            tmp_img.unlink(missing_ok=True)

        save_file(file_id, markdown)
        json_path = save_json(file_id, structured_json)
        content   = markdown
        extra     = {"json_path": str(json_path), "page_count": 1, "block_count": len(structured_json.get("blocks", []))}

    elif ext in DOCX_EXTS:
        tmp_docx = STORAGE_DIR / f"{file_id}_tmp{ext}"
        tmp_docx.write_bytes(raw_content)
        try:
            content = docx_to_markdown(tmp_docx, file_id, filename)
        finally:
            tmp_docx.unlink(missing_ok=True)

        save_file(file_id, content)
        extra = {}

    else:
        content = raw_content.decode("utf-8", errors="replace").strip()
        save_file(file_id, content)
        extra = {}

    save_meta(file_id, {"type": "uploaded", "filename": f"{file_id}.md", "original_filename": filename, "url": None, "created_at": datetime.utcnow().isoformat() + "Z", **extra})

    # Index for RAG
    index_document(file_id, content)

    return {"file_id": file_id, "content": content, "filename": f"{file_id}.md", "original_filename": filename, "sizeBytes": len(raw_content), **extra}


@app.get("/content/{file_id}")
def get_content(file_id: str):
    content = load_file(file_id)
    meta    = load_meta(file_id)
    return {"file_id": file_id, "content": content, "type": meta.get("type", "unknown"), "filename": meta.get("filename", f"{file_id}.md"), "url": meta.get("url")}


@app.get("/content/{file_id}/json")
def get_json_content(file_id: str):
    json_path = STORAGE_DIR / f"{file_id}.json"
    if not json_path.exists():
        raise HTTPException(status_code=404, detail="JSON extraction not found for this file.")
    return json.loads(json_path.read_text(encoding="utf-8"))



@app.get("/files")
def list_files():
    files = []
    for path in STORAGE_DIR.glob("*.md"):
        file_id = path.stem
        meta    = load_meta(file_id)
        # Always display as .md
        raw_name = meta.get("filename") or f"{file_id}.md"
        if not raw_name.endswith(".md"):
            raw_name = Path(raw_name).stem + ".md"
        files.append({
            "file_id":     file_id,
            "filename":    raw_name,
            "type":        meta.get("type", "scraped"),
            "url":         meta.get("url"),
            "size_bytes":  path.stat().st_size,
            "last_modified": path.stat().st_mtime,
            "has_json":    (STORAGE_DIR / f"{file_id}.json").exists(),
            "page_count":  meta.get("page_count"),
            "block_count": meta.get("block_count"),
        })
    files.sort(key=lambda f: f["last_modified"], reverse=True)
    return {"files": files}


@app.delete("/content/{file_id}")
def delete_content(file_id: str):
    md_path = STORAGE_DIR / f"{file_id}.md"
    if not md_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    md_path.unlink()

    for p in [STORAGE_DIR / f"{file_id}.meta.json", STORAGE_DIR / f"{file_id}.json"]:
        if p.exists():
            p.unlink()

    img_dir = IMAGES_DIR / file_id
    if img_dir.exists():
        shutil.rmtree(img_dir)

    # Remove from FAISS index
    remove_document_from_index(file_id)

    return {"file_id": file_id, "message": "Deleted successfully"}


# ─── RAG Routes ─────────────────────────────────────────────────────────────── #

@app.post("/rag/query")
def rag_query(body: RAGQueryRequest):
    """
    Query the RAG system:
    1. Embed the query
    2. Find top_k most relevant chunks from FAISS
    3. Send chunks + query to OpenAI for a grounded answer
    """
    chunks = search_index(body.query, top_k=body.top_k)

    if not chunks:
        return {
            "answer": "No documents have been indexed yet. Please upload or scrape some documents first.",
            "sources": [],
            "chunks_used": [],
        }

    # Build context from retrieved chunks
    context_parts = []
    for chunk in chunks:
        context_parts.append(chunk['text'])

    context = "\n\n---\n\n".join(context_parts)

    # Ask OpenAI
    response = ai.chat.completions.create(
        model="gpt-4o",
        max_tokens=1500,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a helpful assistant. Answer the user's question using ONLY "
                    "the provided context. If the answer is not in the context, say so clearly."
                ),
            },
            {
                "role": "user",
                "content": f"Context:\n{context}\n\nQuestion: {body.query}",
            },
        ],
    )

    answer = response.choices[0].message.content.strip()

    return {"answer": answer}


@app.get("/rag/index/stats")
def rag_index_stats():
    """Return info about the current FAISS index."""
    total_vectors = _faiss_index.ntotal if _faiss_index else 0
    file_ids = list({m["file_id"] for m in _faiss_meta})
    return {
        "total_vectors": total_vectors,
        "total_documents": len(file_ids),
        "indexed_file_ids": file_ids,
    }


@app.post("/rag/index/rebuild")
def rag_rebuild_index():
    """Re-index ALL documents from scratch. Useful after bulk imports."""
    global _faiss_index, _faiss_meta
    _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
    _faiss_meta  = []
    _save_faiss()

    indexed = 0
    for path in STORAGE_DIR.glob("*.md"):
        file_id = path.stem
        content = path.read_text(encoding="utf-8")
        index_document(file_id, content)
        indexed += 1

    return {"message": f"Rebuilt index for {indexed} documents.", "total_vectors": _faiss_index.ntotal}


@app.delete("/rag/index/{file_id}")
def rag_remove_from_index(file_id: str):
    """Manually remove a document from the FAISS index."""
    remove_document_from_index(file_id)
    return {"file_id": file_id, "message": "Removed from index."}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000, reload=True)