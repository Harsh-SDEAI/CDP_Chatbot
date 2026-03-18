"""
Business logic: Admin FAISS indexing, document processing (docling),
storage helpers, KB export, and LlamaIndex chatbot engine.
"""
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import faiss
import numpy as np
from bs4 import BeautifulSoup
from openai import OpenAI

from config import (
    OPENAI_API_KEY, STORAGE_DIR, IMAGES_DIR,
    FAISS_INDEX_PATH, FAISS_META_PATH, EMBEDDING_DIM,
    TEXT_FOLDER, SAVE_PATH, MAX_FILE_SIZE, BASE_FILE_NAME, DATA_DIR,
    get_db_connection, clean_text,
)

# ── Docling imports (PDF / image OCR) ─────────────────────────────────────────
from docling.document_converter import DocumentConverter, PdfFormatOption, ImageFormatOption
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions, EasyOcrOptions

# ── LlamaIndex imports (chatbot engine) ───────────────────────────────────────
from llama_index.core import SimpleDirectoryReader, VectorStoreIndex, StorageContext, load_index_from_storage
from llama_index.vector_stores.faiss import FaissVectorStore
from llama_index.llms.openai import OpenAI as LlamaOpenAI
from llama_index.embeddings.openai import OpenAIEmbedding

# ═══════════════════════════════════════════════════════════════════════════════
#  ADMIN FAISS INDEX  (raw FAISS + text-embedding-3-small, 1536-dim)
# ═══════════════════════════════════════════════════════════════════════════════

ai = OpenAI()  # uses OPENAI_API_KEY from env

_faiss_index: Optional[faiss.IndexFlatL2] = None
_faiss_meta: List[Dict[str, Any]] = []


def load_admin_faiss():
    global _faiss_index, _faiss_meta
    if FAISS_INDEX_PATH.exists() and FAISS_META_PATH.exists():
        _faiss_index = faiss.read_index(str(FAISS_INDEX_PATH))
        _faiss_meta = json.loads(FAISS_META_PATH.read_text(encoding="utf-8"))
        print(f"[FAISS] Loaded index with {_faiss_index.ntotal} vectors.")
    else:
        _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
        _faiss_meta = []
        print("[FAISS] Created new empty index.")


def _save_faiss():
    faiss.write_index(_faiss_index, str(FAISS_INDEX_PATH))
    FAISS_META_PATH.write_text(json.dumps(_faiss_meta, ensure_ascii=False), encoding="utf-8")


def _get_embedding(text: str) -> np.ndarray:
    response = ai.embeddings.create(
        model="text-embedding-3-small",
        input=text[:8000],
    )
    return np.array(response.data[0].embedding, dtype=np.float32)


def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    words = text.split()
    chunks, i = [], 0
    while i < len(words):
        chunks.append(" ".join(words[i:i + chunk_size]))
        i += chunk_size - overlap
    return chunks


def index_document(file_id: str, content: str):
    global _faiss_index, _faiss_meta
    remove_document_from_index(file_id)

    chunks = _chunk_text(content)
    vectors, metas = [], []

    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        try:
            vec = _get_embedding(chunk)
            vectors.append(vec)
            metas.append({"file_id": file_id, "chunk_index": i, "text": chunk})
        except Exception as e:
            print(f"[FAISS] Embedding failed for chunk {i} of {file_id}: {e}")

    if vectors:
        _faiss_index.add(np.stack(vectors))
        _faiss_meta.extend(metas)
        _save_faiss()
        print(f"[FAISS] Indexed {len(vectors)} chunks for {file_id}.")


def remove_document_from_index(file_id: str):
    global _faiss_index, _faiss_meta
    remaining = [m for m in _faiss_meta if m["file_id"] != file_id]
    if len(remaining) == len(_faiss_meta):
        return

    _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
    _faiss_meta = []

    if remaining:
        vectors = []
        for m in remaining:
            try:
                vectors.append(_get_embedding(m["text"]))
                _faiss_meta.append(m)
            except Exception:
                pass
        if vectors:
            _faiss_index.add(np.stack(vectors))

    _save_faiss()
    print(f"[FAISS] Removed chunks for {file_id}, index now has {_faiss_index.ntotal} vectors.")


def search_index(query: str, top_k: int = 5) -> List[Dict[str, Any]]:
    if _faiss_index is None or _faiss_index.ntotal == 0:
        return []

    query_vec = _get_embedding(query)
    k = min(top_k, _faiss_index.ntotal)
    distances, indices = _faiss_index.search(np.array([query_vec], dtype=np.float32), k)

    results = []
    for dist, idx in zip(distances[0], indices[0]):
        if idx == -1:
            continue
        meta = _faiss_meta[idx].copy()
        meta["score"] = float(dist)
        results.append(meta)
    return results


def get_admin_faiss_stats():
    total_vectors = _faiss_index.ntotal if _faiss_index else 0
    file_ids = list({m["file_id"] for m in _faiss_meta})
    return {"total_vectors": total_vectors, "total_documents": len(file_ids), "indexed_file_ids": file_ids}


def rebuild_admin_faiss():
    global _faiss_index, _faiss_meta
    _faiss_index = faiss.IndexFlatL2(EMBEDDING_DIM)
    _faiss_meta = []
    _save_faiss()

    indexed = 0
    for path in STORAGE_DIR.glob("*.md"):
        file_id = path.stem
        content = path.read_text(encoding="utf-8")
        index_document(file_id, content)
        indexed += 1

    return {"message": f"Rebuilt index for {indexed} documents.", "total_vectors": _faiss_index.ntotal}


def rag_query(query: str, top_k: int = 5) -> dict:
    chunks = search_index(query, top_k=top_k)
    if not chunks:
        return {
            "answer": "No documents have been indexed yet. Please upload or scrape some documents first.",
            "sources": [], "chunks_used": [],
        }
    context = "\n\n---\n\n".join(c["text"] for c in chunks)
    response = ai.chat.completions.create(
        model="gpt-4o", max_tokens=1500,
        messages=[
            {"role": "system", "content": "You are a helpful assistant. Answer the user's question using ONLY the provided context. If the answer is not in the context, say so clearly."},
            {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {query}"},
        ],
    )
    return {"answer": response.choices[0].message.content.strip()}


# ═══════════════════════════════════════════════════════════════════════════════
#  STORAGE HELPERS  (markdown files on disk)
# ═══════════════════════════════════════════════════════════════════════════════

def save_file(file_id: str, content: str) -> Path:
    path = STORAGE_DIR / f"{file_id}.md"
    path.write_text(content, encoding="utf-8")
    return path


def save_text_file(file_id: str, content: str) -> Path:
    """Save a plain .txt copy of the extracted text for easy comparison."""
    import re
    # Strip markdown formatting: headings, bold, italic, links, images, code fences
    text = content
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)  # headings
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)       # images
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)        # links
    text = re.sub(r"```[^\n]*\n(.*?)```", r"\1", text, flags=re.DOTALL)  # code fences
    text = re.sub(r"`([^`]+)`", r"\1", text)                    # inline code
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)                # bold
    text = re.sub(r"\*(.+?)\*", r"\1", text)                    # italic
    text = re.sub(r"^[>\-\*]\s+", "", text, flags=re.MULTILINE) # blockquote/list markers
    text = re.sub(r"\n{3,}", "\n\n", text)                      # collapse blank lines
    path = STORAGE_DIR / f"{file_id}.txt"
    path.write_text(text.strip(), encoding="utf-8")
    return path


def load_file(file_id: str) -> str:
    path = STORAGE_DIR / f"{file_id}.md"
    if not path.exists():
        raise FileNotFoundError(f"File {file_id} not found")
    return path.read_text(encoding="utf-8")


def save_meta(file_id: str, meta: dict):
    (STORAGE_DIR / f"{file_id}.meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")


def load_meta(file_id: str) -> dict:
    path = STORAGE_DIR / f"{file_id}.meta.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def save_json(file_id: str, data: dict) -> Path:
    path = STORAGE_DIR / f"{file_id}.json"
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def make_file_id(filename: str) -> str:
    stem = Path(filename).stem
    safe = re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9_\-]", "_", stem)).strip("_")
    candidate, counter = safe, 2
    while (STORAGE_DIR / f"{candidate}.md").exists():
        candidate = f"{safe}_{counter}"
        counter += 1
    return candidate


# ═══════════════════════════════════════════════════════════════════════════════
#  HTML → MARKDOWN
# ═══════════════════════════════════════════════════════════════════════════════

def html_to_markdown(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    lines = []
    for elem in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote"]):
        tag = elem.name
        text = elem.get_text(separator=" ", strip=True)
        if not text:
            continue
        if tag == "h1":           lines.append(f"# {text}")
        elif tag == "h2":         lines.append(f"## {text}")
        elif tag == "h3":         lines.append(f"### {text}")
        elif tag == "h4":         lines.append(f"#### {text}")
        elif tag == "li":         lines.append(f"- {text}")
        elif tag == "pre":        lines.append(f"```\n{text}\n```")
        elif tag == "blockquote": lines.append(f"> {text}")
        else:                     lines.append(text)
        lines.append("")
    return "\n".join(lines).strip()


# ═══════════════════════════════════════════════════════════════════════════════
#  DOCLING HELPERS  (PDF / image OCR extraction)
# ═══════════════════════════════════════════════════════════════════════════════

BULLET_CHARS = ["•", "·", "●", "▪", "", "o", "◦", "-"]
CHECKBOX_CHARS = ["□", "☐", "✓", "✔", "✗", "✘", ""]


def _get_by_ref(doc: Dict[str, Any], ref_str: str) -> Any:
    if not ref_str.startswith("#/"):
        raise ValueError(f"Bad ref: {ref_str}")
    obj: Any = doc
    for p in ref_str[2:].split("/"):
        obj = obj[int(p)] if isinstance(obj, list) else obj[p]
    return obj


def _extract_pages(text_obj: Dict[str, Any]) -> set:
    pages = set()
    for prov in text_obj.get("prov", []):
        if prov.get("page_no") is not None:
            pages.add(int(prov["page_no"]))
    return pages


def _pymupdf_extract_image_text(pdf_path: Optional[Path], obj: Dict[str, Any]) -> str:
    """Fallback: use PyMuPDF to extract text from the bounding box of a picture element."""
    if pdf_path is None:
        return ""
    try:
        import fitz
        prov_list = obj.get("prov", [])
        if not prov_list:
            return ""
        prov = prov_list[0]
        page_no = prov.get("page_no")
        bbox = prov.get("bbox")
        if page_no is None or bbox is None:
            return ""
        doc = fitz.open(str(pdf_path))
        # Docling uses 1-based page numbers
        page = doc[int(page_no) - 1]
        # bbox format from docling: {"l": left, "t": top, "r": right, "b": bottom}
        rect = fitz.Rect(
            float(bbox.get("l", 0)),
            float(bbox.get("t", 0)),
            float(bbox.get("r", 0)),
            float(bbox.get("b", 0)),
        )
        text = page.get_text("text", clip=rect).strip()
        doc.close()
        return text
    except Exception as e:
        print(f"[Docling] PyMuPDF fallback failed for picture: {e}")
        return ""


def _pymupdf_supplementary_text(pdf_path: Path, docling_md: str) -> str:
    """Extract all text from the PDF via PyMuPDF and return lines that Docling missed."""
    try:
        import fitz
        doc = fitz.open(str(pdf_path))
        missed_blocks: List[str] = []
        docling_lower = docling_md.lower()

        for page in doc:
            blocks = page.get_text("blocks")  # list of (x0, y0, x1, y1, text, block_no, type)
            for block in blocks:
                if block[6] != 0:  # type 0 = text, 1 = image
                    continue
                text = block[4].strip()
                if not text or len(text) < 5:
                    continue
                # Check if this text block (or its significant parts) is already in Docling output
                # Split into lines and check each one
                lines = [l.strip() for l in text.split("\n") if l.strip()]
                missing_lines = []
                for line in lines:
                    # A line is "missing" if it doesn't appear in the Docling markdown
                    if line.lower() not in docling_lower:
                        missing_lines.append(line)
                if missing_lines:
                    missed_blocks.append("\n".join(missing_lines))

        doc.close()

        if missed_blocks:
            combined = "\n\n".join(missed_blocks)
            print(f"[Docling] PyMuPDF supplementary extraction found {len(missed_blocks)} missed text block(s)")
            return combined
        return ""
    except Exception as e:
        print(f"[Docling] PyMuPDF supplementary extraction failed: {e}")
        return ""


def _clean_bullet(s: str) -> str:
    s = (s or "").lstrip()
    if not s:
        return s
    for ch in BULLET_CHARS + CHECKBOX_CHARS:
        if s.startswith(ch):
            return s[len(ch):].lstrip()
    return re.sub(r"^[^\w\d]+", "", s).lstrip()


def _format_list_item(obj: Dict[str, Any], indent: int) -> str:
    text = (obj.get("text") or "").strip()
    orig = (obj.get("orig") or "").strip()
    marker = (obj.get("marker") or "").strip()
    pad = "    " * indent
    if obj.get("enumerated") and marker:
        return pad + f"{marker} {text}"
    base = _clean_bullet(orig) or text
    return pad + f"- {base}"


def _format_checkbox(obj: Dict[str, Any], indent: int) -> str:
    orig = (obj.get("orig") or "").strip()
    text = (obj.get("text") or "").strip()
    pad = "    " * indent
    base = _clean_bullet(orig) or _clean_bullet(text) or text
    label = obj.get("label") or ""
    box = "[x]" if "selected" in label else "[ ]"
    return f"{pad}{box} {base}"


def _docling_doc_to_markdown(doc_dict: Dict[str, Any], pdf_path: Optional[Path] = None) -> str:
    lines: List[str] = []

    def _render_table(table_obj):
        grid = table_obj.get("grid") or table_obj.get("data", {}).get("grid", [])
        if not grid:
            return []
        all_texts = [(cell.get("text") or "").strip() for row in grid for cell in row]
        is_menu = any(re.match(r'^\$\d+', t) for t in all_texts)

        if is_menu:
            tlines, seen = [], set()
            for row in grid:
                unique_cells = []
                for cell in row:
                    t = (cell.get("text") or "").strip()
                    if t and t not in seen:
                        seen.add(t)
                        unique_cells.append(t)
                if not unique_cells:
                    continue
                price = next((t for t in unique_cells if re.match(r'^\$\d+', t)), None)
                content_cells = [t for t in unique_cells if not re.match(r'^\$\d+', t)]
                content = " | ".join(content_cells) if content_cells else ""
                if content and price:
                    tlines.append(f"- {content} {price}")
                elif content:
                    tlines.append(f"- {content}")
            return tlines

        tlines = []
        for row_idx, row in enumerate(grid):
            cells = [str((cell.get("text") or "").strip()) for cell in row]
            tlines.append("| " + " | ".join(cells) + " |")
            if row_idx == 0:
                tlines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        return tlines

    def walk_ref(ref: str, indent: int = 0):
        kind = ref.split("/")[1]
        try:
            obj = _get_by_ref(doc_dict, ref)
        except Exception:
            return

        if kind == "texts":
            if obj.get("content_layer") != "body":
                return
            label = obj.get("label", "")
            text = (obj.get("text") or "").strip()
            if not text:
                return
            if label == "section_header":
                lines.append(f"\n{'#' * obj.get('level', 2)} {text}\n")
            elif label == "list_item":
                lines.append(_format_list_item(obj, indent))
            elif label and label.startswith("checkbox"):
                lines.append(_format_checkbox(obj, indent))
            else:
                lines.append(text)

        elif kind == "groups":
            child_indent = indent + 1 if obj.get("label") == "list" else indent
            for child in obj.get("children", []):
                walk_ref(child["$ref"], child_indent)

        elif kind == "tables":
            tlines = _render_table(obj)
            if tlines:
                lines.append("")
                lines.extend(tlines)
                lines.append("")

        elif kind == "pictures":
            captions, ocr_texts = [], []
            for ch in obj.get("children", []):
                try:
                    t = _get_by_ref(doc_dict, ch["$ref"])
                    text = (t.get("text") or "").strip()
                    label = (t.get("label") or "").strip()
                    if not text:
                        continue
                    (captions if label in ("caption", "footnote") else ocr_texts).append(text)
                except Exception as e:
                    print(f"[Docling] Warning: failed to extract picture child: {e}")

            # Fallback: use PyMuPDF to extract text from the picture's bounding box
            if not ocr_texts:
                fallback_text = _pymupdf_extract_image_text(pdf_path, obj)
                if fallback_text:
                    print(f"[Docling] PyMuPDF fallback extracted text from picture region")
                    ocr_texts.append(fallback_text)

            parts = []
            if captions:
                parts.append("Caption: " + " ".join(captions))
            if ocr_texts:
                # Include image text as a proper content block, not just a note
                lines.append("")
                lines.append("**Image Content:**")
                lines.append(" ".join(ocr_texts))
                lines.append("")
            if captions and not ocr_texts:
                lines.append(f"*[Image — {' | '.join(parts)}]*")
            elif not ocr_texts and not captions:
                print(f"[Docling] Warning: picture element produced no text or caption")

    for child in doc_dict.get("body", {}).get("children", []):
        walk_ref(child["$ref"], indent=0)
        lines.append("")

    docling_md = "\n".join(lines).strip()

    # ── Supplementary PyMuPDF extraction ──────────────────────────────
    # Docling can miss text inside boxes, frames, or non-body content layers.
    # Run a full PyMuPDF text extraction and append anything that was missed.
    if pdf_path is not None:
        extra = _pymupdf_supplementary_text(pdf_path, docling_md)
        if extra:
            docling_md += "\n\n" + extra

    return docling_md


def _docling_doc_to_structured_json(doc_dict: Dict[str, Any], file_id: str, filename: str, pdf_path: Optional[Path] = None) -> Dict[str, Any]:
    blocks: List[Dict[str, Any]] = []

    def walk_ref(ref: str, current_section: Optional[str] = None):
        kind = ref.split("/")[1]
        try:
            obj = _get_by_ref(doc_dict, ref)
        except Exception:
            return
        pages = sorted(_extract_pages(obj))

        if kind == "texts":
            if obj.get("content_layer") != "body":
                return
            text = (obj.get("text") or "").strip()
            if not text:
                return
            blocks.append({"type": obj.get("label", "paragraph"), "text": text, "section": current_section, "pages": pages})

        elif kind == "groups":
            items, g_pages = [], set()
            for ch in obj.get("children", []):
                try:
                    t = _get_by_ref(doc_dict, ch["$ref"])
                    text = (t.get("text") or "").strip()
                    if text:
                        items.append(text)
                        g_pages.update(_extract_pages(t))
                except Exception:
                    pass
            if items:
                blocks.append({"type": obj.get("label", "group"), "items": items, "section": current_section, "pages": sorted(g_pages)})

        elif kind == "tables":
            grid = obj.get("grid", [])
            rows = [[str((c.get("text") or "").strip()) for c in row] for row in grid]
            if rows:
                blocks.append({"type": "table", "headers": rows[0], "rows": rows[1:], "section": current_section, "pages": pages})

        elif kind == "pictures":
            captions, ocr_texts, p_pages = [], [], set(pages)
            for ch in obj.get("children", []):
                try:
                    t = _get_by_ref(doc_dict, ch["$ref"])
                    text = (t.get("text") or "").strip()
                    label = (t.get("label") or "").strip()
                    if not text:
                        continue
                    p_pages.update(_extract_pages(t))
                    (captions if label in ("caption", "footnote") else ocr_texts).append(text)
                except Exception as e:
                    print(f"[Docling] Warning: failed to extract picture child in JSON: {e}")

            # Fallback: use PyMuPDF to extract text from the picture's bounding box
            if not ocr_texts:
                fallback_text = _pymupdf_extract_image_text(pdf_path, obj)
                if fallback_text:
                    ocr_texts.append(fallback_text)

            if captions or ocr_texts:
                blocks.append({"type": "picture", "captions": captions, "ocr_text": ocr_texts, "section": current_section, "pages": sorted(p_pages)})

    current_section = None
    for child in doc_dict.get("body", {}).get("children", []):
        ref = child["$ref"]
        if ref.split("/")[1] == "texts":
            try:
                obj = _get_by_ref(doc_dict, ref)
                if obj.get("label") == "section_header":
                    current_section = (obj.get("text") or "").strip()
            except Exception:
                pass
        walk_ref(ref, current_section)

    origin = doc_dict.get("origin", {})
    return {
        "file_id": file_id, "filename": filename,
        "page_count": len(doc_dict.get("pages", {})),
        "extracted_at": datetime.utcnow().isoformat() + "Z",
        "source": {"mime_type": origin.get("mimetype"), "original_filename": origin.get("filename")},
        "blocks": blocks,
    }


def image_to_docling(img_path: Path, file_id: str, filename: str):
    print(f"[Docling] Starting image extraction for {file_id}...")
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = True
    pipeline_options.ocr_options = EasyOcrOptions(lang=["en"])
    pipeline_options.do_table_structure = True
    pipeline_options.images_scale = 2.0

    converter = DocumentConverter(
        format_options={InputFormat.IMAGE: ImageFormatOption(pipeline_options=pipeline_options)}
    )
    result = converter.convert(str(img_path))
    doc_dict = result.document.export_to_dict()
    print(f"[Docling] Image extraction complete for {file_id}.")

    markdown = _docling_doc_to_markdown(doc_dict)
    structured_json = _docling_doc_to_structured_json(doc_dict, file_id, filename)
    structured_json["docling_raw"] = doc_dict
    return markdown, structured_json


def docx_to_text(docx_path: Path, file_id: str, filename: str):
    """Extract text from a Word (.docx) file using python-docx."""
    from docx import Document as DocxDocument

    print(f"[DOCX] Starting text extraction for {file_id}...")
    doc = DocxDocument(str(docx_path))

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                paragraphs.append(row_text)

    markdown = "\n\n".join(paragraphs)
    print(f"[DOCX] Extraction complete for {file_id} ({len(paragraphs)} blocks).")
    return markdown


_PDF_CHUNK_SIZE = 10  # pages per batch – keeps peak RAM in check


def _build_pdf_pipeline():
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = True
    pipeline_options.ocr_options = EasyOcrOptions(lang=["en"])
    pipeline_options.do_table_structure = True
    pipeline_options.images_scale = 2.0
    pipeline_options.generate_picture_images = True
    return pipeline_options


def pdf_to_docling(pdf_path: Path, file_id: str, filename: str):
    import fitz  # PyMuPDF

    doc = fitz.open(str(pdf_path))
    total_pages = len(doc)
    doc.close()
    print(f"[Docling] Starting OCR extraction for {file_id} ({total_pages} pages)...")

    if total_pages <= _PDF_CHUNK_SIZE:
        # Small PDF → single pass (original behaviour)
        return _convert_pdf_single(pdf_path, file_id, filename)

    # Large PDF → process in chunks of _PDF_CHUNK_SIZE pages
    return _convert_pdf_chunked(pdf_path, file_id, filename, total_pages)


def _convert_pdf_single(pdf_path: Path, file_id: str, filename: str):
    pipeline_options = _build_pdf_pipeline()
    converter = DocumentConverter(
        format_options={"pdf": PdfFormatOption(pipeline_options=pipeline_options)}
    )
    result = converter.convert(str(pdf_path))
    doc_dict = result.document.export_to_dict()
    print(f"[Docling] Extraction complete for {file_id}.")

    markdown = _docling_doc_to_markdown(doc_dict, pdf_path=pdf_path)
    structured_json = _docling_doc_to_structured_json(doc_dict, file_id, filename, pdf_path=pdf_path)
    structured_json["docling_raw"] = doc_dict
    return markdown, structured_json


def _convert_pdf_chunked(pdf_path: Path, file_id: str, filename: str, total_pages: int):
    """Split a large PDF into temporary chunks, convert each, then merge."""
    import tempfile, gc
    import fitz

    all_md_parts = []
    all_doc_dicts = []

    for start in range(0, total_pages, _PDF_CHUNK_SIZE):
        end = min(start + _PDF_CHUNK_SIZE, total_pages)
        print(f"[Docling]   Processing pages {start + 1}-{end} of {total_pages}...")

        # Extract page range into a temp PDF
        src = fitz.open(str(pdf_path))
        tmp_pdf = fitz.open()
        tmp_pdf.insert_pdf(src, from_page=start, to_page=end - 1)

        # On Windows, NamedTemporaryFile keeps the handle open → permission error.
        # Create the path, close the handle, then save.
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".pdf")
        os.close(tmp_fd)
        tmp_pdf.save(tmp_path)
        tmp_pdf.close()
        src.close()

        pipeline_options = _build_pdf_pipeline()
        converter = DocumentConverter(
            format_options={"pdf": PdfFormatOption(pipeline_options=pipeline_options)}
        )
        result = converter.convert(tmp_path)
        doc_dict = result.document.export_to_dict()

        all_md_parts.append(_docling_doc_to_markdown(doc_dict, pdf_path=Path(tmp_path)))
        all_doc_dicts.append(doc_dict)

        del converter, result
        gc.collect()
        try:
            os.remove(tmp_path)
        except OSError:
            pass  # cleanup is best-effort

    print(f"[Docling] Extraction complete for {file_id}.")

    markdown = "\n\n".join(all_md_parts)
    # Use the first chunk's dict as the base for structured JSON
    structured_json = _docling_doc_to_structured_json(all_doc_dicts[0], file_id, filename, pdf_path=pdf_path)
    structured_json["docling_raw"] = all_doc_dicts
    return markdown, structured_json


# ═══════════════════════════════════════════════════════════════════════════════
#  KB EXPORT  (database → text files for LlamaIndex)
# ═══════════════════════════════════════════════════════════════════════════════

def get_last_file_index_and_size():
    if not os.path.exists(SAVE_PATH):
        os.makedirs(SAVE_PATH)
        return 1, 0
    files = [f for f in os.listdir(SAVE_PATH) if re.match(rf"{BASE_FILE_NAME}(\d+)\.txt", f)]
    if not files:
        return 1, 0
    files.sort(key=lambda x: int(re.findall(rf"{BASE_FILE_NAME}(\d+)\.txt", x)[0]))
    last_file = files[-1]
    last_index = int(re.findall(rf"{BASE_FILE_NAME}(\d+)\.txt", last_file)[0])
    size = os.path.getsize(os.path.join(SAVE_PATH, last_file))
    return last_index, size


def export_qa_pairs_job():
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute(
        "SELECT id, Question, Answer, Status, SuggestedAnswer "
        "FROM CDPChatHistory WHERE IsKBUpdated = TRUE AND PublishStatus = 'ReadyToPublish'"
    )
    rows = cursor.fetchall()
    if not rows:
        conn.close()
        return {"status": "ok", "message": "No new records to export", "files": []}

    file_index, last_file_size = get_last_file_index_and_size()
    current_file_path = os.path.join(SAVE_PATH, f"{BASE_FILE_NAME}{file_index}.txt")
    mode = 'a' if last_file_size < MAX_FILE_SIZE else 'w'
    if mode == 'w':
        file_index += 1
        current_file_path = os.path.join(SAVE_PATH, f"{BASE_FILE_NAME}{file_index}.txt")

    current_file = open(current_file_path, mode, encoding='utf-8')
    created_files = {os.path.basename(current_file_path)}

    for row in rows:
        chat_id = row[0]
        print(f"Processing ChatHistoryID: {chat_id}")
        question = row[1] or ''
        answer = clean_text(row[2])
        suggested = clean_text(row[4])
        status = (row[3] or '').strip()

        if status == 'PartiallyCorrect':
            info = f"{answer.strip()} {suggested.strip()}"
        elif status == 'InCorrect':
            info = suggested.strip()
        else:
            continue

        pair_text = f"Title: {question.strip()}\nInformation: {info}\n\n"

        if current_file.tell() + len(pair_text.encode('utf-8')) > MAX_FILE_SIZE:
            current_file.close()
            file_index += 1
            current_file_path = os.path.join(SAVE_PATH, f"{BASE_FILE_NAME}{file_index}.txt")
            current_file = open(current_file_path, 'w', encoding='utf-8')
            created_files.add(os.path.basename(current_file_path))

        current_file.write(pair_text)
        cursor.execute(
            "UPDATE CDPChatHistory SET IsKBUpdated = FALSE, PublishStatus = 'Published' WHERE id = %s",
            (chat_id,),
        )
    conn.commit()
    current_file.close()
    conn.close()
    return {"status": "ok", "files": sorted(list(created_files))}


# ═══════════════════════════════════════════════════════════════════════════════
#  LLAMAINDEX CHATBOT ENGINE  (FAISS + text-embedding-3-large, 3072-dim)
# ═══════════════════════════════════════════════════════════════════════════════

openai_llm = LlamaOpenAI(
    api_key=OPENAI_API_KEY,
    model_name="gpt-3.5-turbo",
    temperature=0.1,
    top_p=0.8,
    max_tokens=1024,
)

# Global query engine and embed model — loaded on startup, rebuilt via /update-index
query_engine = None
chatbot_embed_model = OpenAIEmbedding(model="text-embedding-3-large", api_key=OPENAI_API_KEY)


def load_chatbot_index():
    """Load existing LlamaIndex FAISS index from disk (called at startup)."""
    global query_engine
    persist_dir = "./storage"
    try:
        vector_store = FaissVectorStore.from_persist_dir(persist_dir)
        storage_context = StorageContext.from_defaults(vector_store=vector_store, persist_dir=persist_dir)
        index = load_index_from_storage(storage_context=storage_context, embed_model=chatbot_embed_model)
        query_engine = index.as_query_engine(llm=openai_llm, similarity_top_k=7)
        print("Loaded chatbot FAISS index from local storage.")
    except Exception as e:
        print(f"[Chatbot] Could not load index: {e}. Run /update-index to build one.")


def build_chatbot_faiss_index():
    """Rebuild the LlamaIndex FAISS index from TEXT_FOLDER and Data documents."""
    global query_engine

    documents = []
    supported_exts = [".txt", ".md"]

    # Load from storage folder (KB exports + uploaded docs)
    try:
        docs = SimpleDirectoryReader(TEXT_FOLDER, required_exts=supported_exts).load_data()
        documents.extend(docs)
        print(f"Loaded {len(docs)} documents from {TEXT_FOLDER}")
    except ValueError:
        print(f"No supported files found in {TEXT_FOLDER}")

    # Also load from Data folder if it exists (base knowledge)
    data_dir = str(DATA_DIR)
    try:
        if DATA_DIR.exists():
            docs = SimpleDirectoryReader(data_dir, required_exts=supported_exts).load_data()
            documents.extend(docs)
            print(f"Loaded {len(docs)} documents from {data_dir}")
    except ValueError:
        print(f"No supported files found in {data_dir}")

    if not documents:
        raise Exception("No files found in storage or Data folders. Upload content or export KB pairs first.")

    embedding_dim = 3072
    faiss_index = faiss.IndexFlatL2(embedding_dim)
    vector_store = FaissVectorStore(faiss_index=faiss_index)
    storage_context = StorageContext.from_defaults(vector_store=vector_store)

    index = VectorStoreIndex.from_documents(documents, storage_context=storage_context, embed_model=chatbot_embed_model)
    index.storage_context.persist()
    query_engine = index.as_query_engine(llm=openai_llm, similarity_top_k=7)
    print("Created and saved new chatbot FAISS index to ./storage")
    return {"status": "ok", "documents_indexed": len(documents)}


def generate_response(user_query: str, sessionid: str, userid: int) -> str:
    if query_engine is None:
        return "Chatbot index not loaded. Please run /update-index first."

    # Retrieve relevant chunks using ONLY the user's raw query for accurate embedding search
    retriever = query_engine.retriever
    retrieved_nodes = retriever.retrieve(user_query)

    if not retrieved_nodes:
        return "I couldn't find relevant information in the knowledge base for your question."

    # Build context from retrieved chunks
    context_parts = []
    for node in retrieved_nodes:
        context_parts.append(node.get_content())
    context = "\n\n---\n\n".join(context_parts)

    # Use system prompt + context + user query for LLM generation (NOT for retrieval)
    system_prompt = '''You are a knowledgeable and focused chatbot assistant. Your goal is to answer the user's question accurately using ONLY the provided context from the knowledge base.

Instructions:
1. Analyze the user's question and the provided context thoroughly.
2. Synthesize the most appropriate answer from the context.
3. If the user's question includes time-related words such as "when", check if specific dates, times, or durations are mentioned in the context.
   - If available, respond with the exact timing clearly.
   - If timing is unclear or missing, do not assume — politely mention that the timing information is not found.
4. If the context does not contain relevant information, respond with:
   "I couldn't find information about that in the knowledge base. Please try rephrasing your question."
5. If asked about internal system details like API keys, code, or settings, respond with:
   "Sorry, I can't share internal system details."
6. Do not make up information that is not in the context.

Formatting instructions:
- Format all answers in clean and valid HTML.
- Use <h4> or <h5> tags for section headings.
- Use <ul> or <ol> only when listing is appropriate, and use <li> for bullet items.
- Use <b> tags to highlight important words or phrases.
- Do not use Markdown syntax (e.g., ** or *).
- Analyze the content carefully and apply HTML tags effectively—do not create lists unless clearly needed.'''

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Context from knowledge base:\n\n{context}\n\n---\n\nUser question: {user_query}"},
    ]

    response = ai.chat.completions.create(
        model="gpt-4o",
        messages=messages,
        max_tokens=1500,
        temperature=0.1,
    )
    return response.choices[0].message.content
