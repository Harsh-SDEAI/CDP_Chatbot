"""
Generate a professional PPT comparing Old Approach (QA Pair Review)
vs New Approach (Content Upload → Markdown → RAG Pipeline).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Color Palette ──
BG_DARK    = RGBColor(0x1A, 0x1A, 0x2E)
BG_CARD    = RGBColor(0x24, 0x24, 0x3E)
ACCENT     = RGBColor(0x4E, 0x8E, 0xFF)
ACCENT2    = RGBColor(0x00, 0xC9, 0xA7)
ORANGE     = RGBColor(0xFF, 0x8C, 0x42)
RED_SOFT   = RGBColor(0xFF, 0x6B, 0x6B)
GREEN      = RGBColor(0x4E, 0xCB, 0x71)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xBB, 0xBB, 0xCC)
MID_GRAY   = RGBColor(0x88, 0x88, 0xAA)
YELLOW     = RGBColor(0xFF, 0xD9, 0x3D)


def add_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color, corner_radius=Emu(120000)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.adjustments[0] = 0.05
    return shape


def add_text(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=14, color=LIGHT_GRAY, bullet_color=ACCENT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        # Bullet character
        run_b = p.add_run()
        run_b.text = "  >  " if bullet_color == ACCENT else "  >  "
        run_b.font.size = Pt(font_size)
        run_b.font.color.rgb = bullet_color
        run_b.font.bold = True
        # Text
        run_t = p.add_run()
        run_t.text = item
        run_t.font.size = Pt(font_size)
        run_t.font.color.rgb = color
    return txBox


def add_numbered_list(slide, left, top, width, height, items, font_size=14, color=LIGHT_GRAY, num_color=ACCENT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(8)
        p.space_before = Pt(2)
        run_n = p.add_run()
        run_n.text = f"  {i+1}.  "
        run_n.font.size = Pt(font_size + 2)
        run_n.font.color.rgb = num_color
        run_n.font.bold = True
        run_t = p.add_run()
        run_t.text = item
        run_t.font.size = Pt(font_size)
        run_t.font.color.rgb = color
    return txBox


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide)

# Accent line
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)

add_text(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(1),
         "CDP CHATBOT", 22, MID_GRAY, bold=True, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(1.5),
         "Content Management Pipeline", 44, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(3.5), Inches(10), Inches(0.8),
         "From QA Pair Review to Intelligent Document Processing & RAG",
         20, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Divider
add_shape(slide, Inches(5.5), Inches(4.5), Inches(2.333), Inches(0.04), ACCENT)

add_text(slide, Inches(2), Inches(5.0), Inches(9), Inches(0.5),
         "Old Approach  vs  New Approach", 20, ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(2), Inches(5.6), Inches(9), Inches(0.7),
         "Cooperstown Dreams Park  |  CDPGPT  |  Content Manager",
         16, MID_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Problem Statement / Why Change?
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "Why the Change?", 36, WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
         "Limitations of the old approach and the need for a better pipeline",
         16, MID_GRAY)

# Left card — Problems
add_shape(slide, Inches(0.6), Inches(1.9), Inches(5.8), Inches(4.8), BG_CARD)
add_text(slide, Inches(1.0), Inches(2.1), Inches(5), Inches(0.5),
         "Challenges with Old Approach", 20, RED_SOFT, bold=True)

add_bullet_list(slide, Inches(0.9), Inches(2.7), Inches(5.2), Inches(4),
    [
        "Manual QA pair review is time-consuming",
        "Admin must review each Q&A individually",
        "Status-based workflow (Correct / Incorrect / Junk) is rigid",
        "Suggestions are per-answer, no bulk content updates",
        "No support for ingesting documents, PDFs, or web content",
        "Knowledge base grows slowly — one QA at a time",
    ], 14, LIGHT_GRAY, RED_SOFT)

# Right card — Vision
add_shape(slide, Inches(6.9), Inches(1.9), Inches(5.8), Inches(4.8), BG_CARD)
add_text(slide, Inches(7.3), Inches(2.1), Inches(5), Inches(0.5),
         "What We Needed", 20, GREEN, bold=True)

add_bullet_list(slide, Inches(7.2), Inches(2.7), Inches(5.2), Inches(4),
    [
        "Upload entire documents at once (PDF, TXT, Images)",
        "Scrape website pages and auto-convert to knowledge",
        "Unified markdown format for easy editing by admin",
        "Automatic vector indexing for instant RAG retrieval",
        "Scalable — add hundreds of pages in minutes",
        "Admin edits content directly, changes reflect in chatbot",
    ], 14, LIGHT_GRAY, GREEN)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Old Approach (Detailed)
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ORANGE)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "Old Approach — QA Pair Review System", 36, WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
         "Manual review of chatbot-generated Q&A pairs through the Config Admin panel",
         16, MID_GRAY)

# Flow steps
add_shape(slide, Inches(0.6), Inches(1.9), Inches(12.1), Inches(1.6), BG_CARD)
add_text(slide, Inches(1.0), Inches(2.0), Inches(11), Inches(0.5),
         "Workflow", 18, ORANGE, bold=True)
add_numbered_list(slide, Inches(0.9), Inches(2.45), Inches(11.5), Inches(1.2),
    [
        "User asks a question to the chatbot (CDPGPT)",
        "Chatbot generates an answer from existing knowledge base",
        "Admin reviews the QA pair in Conversation History (Config Admin panel)",
        "Admin sets a Status: Correct / Incorrect / Partially Correct / Junk",
        "If Incorrect or Partially Correct — admin provides a Suggestion (corrected answer)",
        "Updated answer is saved back to the database for future use",
    ], 13, LIGHT_GRAY, ORANGE)

# Bottom cards
add_shape(slide, Inches(0.6), Inches(3.8), Inches(3.8), Inches(3.1), BG_CARD)
add_text(slide, Inches(1.0), Inches(3.95), Inches(3.4), Inches(0.5),
         "Status Options", 18, YELLOW, bold=True)
add_bullet_list(slide, Inches(0.9), Inches(4.45), Inches(3.4), Inches(2.5),
    [
        "Correct — answer is verified, no change",
        "Incorrect — answer is wrong, admin provides fix",
        "Partially Correct — needs minor corrections",
        "Junk — irrelevant, flagged for removal",
    ], 13, LIGHT_GRAY, YELLOW)

add_shape(slide, Inches(4.8), Inches(3.8), Inches(3.8), Inches(3.1), BG_CARD)
add_text(slide, Inches(5.2), Inches(3.95), Inches(3.4), Inches(0.5),
         "Key Components", 18, ACCENT, bold=True)
add_bullet_list(slide, Inches(5.1), Inches(4.45), Inches(3.4), Inches(2.5),
    [
        "CDPGPT Conversation History page",
        "Filter by Year & Tournament",
        "Action button to open QA review",
        "Save button persists to database",
    ], 13, LIGHT_GRAY, ACCENT)

add_shape(slide, Inches(9.0), Inches(3.8), Inches(3.7), Inches(3.1), BG_CARD)
add_text(slide, Inches(9.4), Inches(3.95), Inches(3.3), Inches(0.5),
         "Limitations", 18, RED_SOFT, bold=True)
add_bullet_list(slide, Inches(9.3), Inches(4.45), Inches(3.3), Inches(2.5),
    [
        "One QA pair at a time",
        "Reactive — only fixes after user asks",
        "No document/URL ingestion",
        "Slow knowledge base growth",
    ], 13, LIGHT_GRAY, RED_SOFT)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — New Approach (Detailed)
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT2)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "New Approach — Document Upload & RAG Pipeline", 36, WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
         "Upload files or scrape URLs → auto-convert to Markdown → index in FAISS → RAG Chat",
         16, MID_GRAY)

# Flow
add_shape(slide, Inches(0.6), Inches(1.9), Inches(12.1), Inches(1.8), BG_CARD)
add_text(slide, Inches(1.0), Inches(2.0), Inches(11), Inches(0.5),
         "Pipeline Flow", 18, ACCENT2, bold=True)
add_numbered_list(slide, Inches(0.9), Inches(2.45), Inches(11.5), Inches(1.5),
    [
        "Admin clicks Upload → chooses File Upload or URL",
        "File: PDF / TXT / MD / Images are processed via Docling (OCR + structure extraction)",
        "URL: Page is scraped with httpx, HTML converted to clean Markdown",
        "Output saved as .md file with descriptive name (e.g. cooperstowndreamspark_com_testimonials.md)",
        "Content is chunked (500 words, 50 overlap) → embedded via OpenAI text-embedding-3-small",
        "Vectors stored in FAISS index → instantly searchable via RAG Chat",
        "Admin can view, edit, or delete any .md file from the Files page",
    ], 12, LIGHT_GRAY, ACCENT2)

# Bottom cards
add_shape(slide, Inches(0.6), Inches(4.0), Inches(3.8), Inches(3.0), BG_CARD)
add_text(slide, Inches(1.0), Inches(4.15), Inches(3.4), Inches(0.5),
         "Input Sources", 18, ACCENT, bold=True)
add_bullet_list(slide, Inches(0.9), Inches(4.6), Inches(3.4), Inches(2.5),
    [
        "PDF documents (OCR via Docling)",
        "Images: PNG, JPG, WEBP, BMP, TIFF",
        "Plain text (.txt) and Markdown (.md)",
        "Any website URL (auto-scrape)",
    ], 13, LIGHT_GRAY, ACCENT)

add_shape(slide, Inches(4.8), Inches(4.0), Inches(3.8), Inches(3.0), BG_CARD)
add_text(slide, Inches(5.2), Inches(4.15), Inches(3.4), Inches(0.5),
         "Processing", 18, ACCENT2, bold=True)
add_bullet_list(slide, Inches(5.1), Inches(4.6), Inches(3.4), Inches(2.5),
    [
        "Docling extracts text, tables, images",
        "HTML → Markdown via BeautifulSoup",
        "Smart file naming from URL/filename",
        "Structured JSON + Markdown output",
    ], 13, LIGHT_GRAY, ACCENT2)

add_shape(slide, Inches(9.0), Inches(4.0), Inches(3.7), Inches(3.0), BG_CARD)
add_text(slide, Inches(9.4), Inches(4.15), Inches(3.3), Inches(0.5),
         "Storage & Retrieval", 18, GREEN, bold=True)
add_bullet_list(slide, Inches(9.3), Inches(4.6), Inches(3.3), Inches(2.5),
    [
        "FAISS vector index (L2 similarity)",
        "OpenAI text-embedding-3-small",
        "GPT-4o for answer generation",
        "Top-k configurable (1–20 chunks)",
    ], 13, LIGHT_GRAY, GREEN)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Side-by-Side Comparison
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "Old vs New — Comparison", 36, WHITE, bold=True)

# Headers
add_shape(slide, Inches(0.6), Inches(1.5), Inches(3.5), Inches(0.6), ORANGE)
add_text(slide, Inches(0.6), Inches(1.53), Inches(3.5), Inches(0.6),
         "Feature", 16, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_shape(slide, Inches(4.3), Inches(1.5), Inches(4.2), Inches(0.6), RGBColor(0x8B, 0x5C, 0x2A))
add_text(slide, Inches(4.3), Inches(1.53), Inches(4.2), Inches(0.6),
         "Old Approach (QA Review)", 16, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_shape(slide, Inches(8.7), Inches(1.5), Inches(4.1), Inches(0.6), RGBColor(0x1A, 0x6B, 0x5A))
add_text(slide, Inches(8.7), Inches(1.53), Inches(4.1), Inches(0.6),
         "New Approach (Content Manager)", 16, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Table rows
rows = [
    ("Content Input",         "Chat-generated QA pairs only",          "PDF, TXT, Images, URLs"),
    ("Admin Workflow",        "Review one QA pair at a time",          "Upload bulk files or scrape pages"),
    ("Data Format",           "QA pairs in database",                  "Markdown (.md) files"),
    ("Editing",               "Status + Suggestion per answer",        "Edit full .md document directly"),
    ("Knowledge Growth",      "Slow — one answer at a time",          "Fast — entire documents at once"),
    ("Search / Retrieval",    "Database keyword lookup",               "FAISS vector similarity (RAG)"),
    ("AI Model",              "GPT-based answer generation",           "GPT-4o + text-embedding-3-small"),
    ("Scalability",           "Limited by QA volume",                  "Upload 100s of pages in minutes"),
    ("Content Sources",       "Only chatbot conversations",            "Any document or website"),
]

y = Inches(2.2)
row_h = Inches(0.52)
for i, (feat, old, new) in enumerate(rows):
    bg = BG_CARD if i % 2 == 0 else RGBColor(0x1E, 0x1E, 0x34)
    add_shape(slide, Inches(0.6), y, Inches(3.5), row_h, bg)
    add_text(slide, Inches(0.8), y + Emu(20000), Inches(3.2), row_h,
             feat, 13, YELLOW, bold=True)

    add_shape(slide, Inches(4.3), y, Inches(4.2), row_h, bg)
    add_text(slide, Inches(4.5), y + Emu(20000), Inches(3.8), row_h,
             old, 12, LIGHT_GRAY)

    add_shape(slide, Inches(8.7), y, Inches(4.1), row_h, bg)
    add_text(slide, Inches(8.9), y + Emu(20000), Inches(3.7), row_h,
             new, 12, ACCENT2, bold=True)

    y += row_h


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Technical Architecture
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "New Approach — Technical Architecture", 36, WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
         "Full stack: Streamlit frontend  +  FastAPI backend  +  FAISS vector store",
         16, MID_GRAY)

# Tech stack boxes
boxes = [
    ("Frontend\nStreamlit", "Content Manager UI\nFile Upload / URL Scrape\nRAG Chat Interface\nFiles Browser & Editor", ACCENT, Inches(0.6)),
    ("Backend\nFastAPI", "POST /upload — file processing\nPOST /scrape — URL fetching\nPOST /rag/query — RAG search\nGET /files — list all docs", ACCENT2, Inches(3.4)),
    ("Processing\nDocling + BS4", "PDF OCR extraction\nImage text recognition\nHTML → Markdown conversion\nTable & structure parsing", ORANGE, Inches(6.2)),
    ("Vector Store\nFAISS + OpenAI", "text-embedding-3-small\nChunking (500w / 50 overlap)\nL2 similarity search\nGPT-4o answer generation", GREEN, Inches(9.0)),
]

for title, desc, color, left in boxes:
    add_shape(slide, left, Inches(1.9), Inches(2.5), Inches(0.8), color)
    add_text(slide, left, Inches(1.92), Inches(2.5), Inches(0.8),
             title, 13, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_shape(slide, left, Inches(2.8), Inches(2.5), Inches(2.5), BG_CARD)
    lines = desc.split("\n")
    add_bullet_list(slide, left + Inches(0.1), Inches(2.9), Inches(2.3), Inches(2.2),
                    lines, 11, LIGHT_GRAY, color)

# Arrows between boxes
for x in [Inches(3.15), Inches(5.95), Inches(8.75)]:
    add_text(slide, x, Inches(2.0), Inches(0.3), Inches(0.5),
             "->", 22, MID_GRAY, bold=True, alignment=PP_ALIGN.CENTER)

# Data flow
add_shape(slide, Inches(0.6), Inches(5.6), Inches(10.8), Inches(1.5), BG_CARD)
add_text(slide, Inches(1.0), Inches(5.7), Inches(10), Inches(0.5),
         "Data Flow Example", 18, YELLOW, bold=True)
add_text(slide, Inches(1.0), Inches(6.15), Inches(10.4), Inches(1.0),
         "URL: cooperstowndreamspark.com/testimonials\n"
         "  -->  Fetch HTML with httpx  -->  Strip nav/footer/scripts  -->  Convert to Markdown\n"
         "  -->  Save as cooperstowndreamspark_com_testimonials.md  -->  Chunk into 500-word segments\n"
         "  -->  Embed each chunk via OpenAI  -->  Add to FAISS index  -->  Ready for RAG queries",
         13, LIGHT_GRAY)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — File Naming Convention
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT2)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "Smart File Naming Convention", 36, WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
         "All content is stored as .md (Markdown) with human-readable names",
         16, MID_GRAY)

# URL examples
add_shape(slide, Inches(0.6), Inches(1.9), Inches(12.1), Inches(1.8), BG_CARD)
add_text(slide, Inches(1.0), Inches(2.0), Inches(5), Inches(0.5),
         "URL Scraping — Name derived from URL path", 18, ACCENT2, bold=True)

url_examples = [
    ("cooperstowndreamspark.com/testimonials/", "cooperstowndreamspark_com_testimonials.md"),
    ("example.com", "example_com.md"),
    ("xyz.com/abc/def", "xyz_com_abc_def.md"),
    ("xyz.com/page.html", "xyz_com_page.md"),
]

y_inner = Inches(2.55)
for url, result in url_examples:
    add_text(slide, Inches(1.2), y_inner, Inches(5), Inches(0.35),
             url, 13, MID_GRAY)
    add_text(slide, Inches(6.5), y_inner, Inches(1), Inches(0.35),
             "-->", 14, ACCENT2, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, Inches(7.5), y_inner, Inches(4.5), Inches(0.35),
             result, 14, ACCENT2, bold=True)
    y_inner += Inches(0.3)

# File upload examples
add_shape(slide, Inches(0.6), Inches(4.0), Inches(12.1), Inches(1.5), BG_CARD)
add_text(slide, Inches(1.0), Inches(4.1), Inches(5), Inches(0.5),
         "File Upload — Name derived from original filename", 18, ACCENT, bold=True)

file_examples = [
    ("img 1.jpg", "img_1.md"),
    ("Annual Report 2025.pdf", "Annual_Report_2025.md"),
    ("sample_test_document.md", "sample_test_document.md"),
    ("photo.png", "photo.md"),
]

y_inner = Inches(4.6)
for orig, result in file_examples:
    add_text(slide, Inches(1.2), y_inner, Inches(5), Inches(0.35),
             orig, 13, MID_GRAY)
    add_text(slide, Inches(6.5), y_inner, Inches(1), Inches(0.35),
             "-->", 14, ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, Inches(7.5), y_inner, Inches(4.5), Inches(0.35),
             result, 14, ACCENT, bold=True)
    y_inner += Inches(0.3)

# Key rule
add_shape(slide, Inches(0.6), Inches(5.8), Inches(12.1), Inches(0.7), RGBColor(0x1A, 0x3A, 0x2E))
add_text(slide, Inches(1.0), Inches(5.85), Inches(11.5), Inches(0.6),
         "Rule: All files are always stored and displayed as .md — the user never sees the original format. "
         "Conversion happens in the background automatically.",
         14, GREEN, bold=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Key Benefits & Summary
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), GREEN)

add_text(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.8),
         "Key Benefits of the New Approach", 36, WHITE, bold=True)

benefits = [
    ("10x Faster Content Ingestion", "Upload entire documents or scrape full websites instead of reviewing one QA pair at a time"),
    ("Unified Markdown Format", "Everything stored as .md — easy to read, edit, and version control. Admin has full control over content"),
    ("Automatic Vector Indexing", "Content is chunked and embedded automatically via FAISS + OpenAI. Instantly searchable via RAG"),
    ("Multi-Source Support", "PDF, TXT, Images (OCR via Docling), Markdown files, and any website URL"),
    ("Smart Naming", "Files named from URLs (cooperstowndreamspark_com_testimonials.md) or original filenames — always .md"),
    ("RAG-Powered Chat", "GPT-4o generates answers grounded in your indexed documents with configurable top-k retrieval"),
]

y = Inches(1.5)
colors = [ACCENT, ACCENT2, GREEN, ORANGE, YELLOW, ACCENT]
for i, (title, desc) in enumerate(benefits):
    add_shape(slide, Inches(0.6), y, Inches(12.1), Inches(0.9), BG_CARD)
    # Number circle
    add_text(slide, Inches(0.8), y + Emu(50000), Inches(0.5), Inches(0.5),
             f"0{i+1}", 16, colors[i], bold=True)
    add_text(slide, Inches(1.5), y + Emu(30000), Inches(3.5), Inches(0.5),
             title, 16, WHITE, bold=True)
    add_text(slide, Inches(5.2), y + Emu(50000), Inches(7.2), Inches(0.5),
             desc, 12, LIGHT_GRAY)
    y += Inches(0.95)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — Thank You
# ═══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)

add_text(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1),
         "Thank You", 48, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_shape(slide, Inches(5.5), Inches(3.2), Inches(2.333), Inches(0.04), ACCENT)

add_text(slide, Inches(2), Inches(3.6), Inches(9), Inches(0.8),
         "CDP Chatbot — Content Management Pipeline",
         20, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(2), Inches(4.4), Inches(9), Inches(0.5),
         "Cooperstown Dreams Park  |  Powered by FAISS + GPT-4o + Docling",
         16, MID_GRAY, alignment=PP_ALIGN.CENTER)

# Tech badges
techs = ["FastAPI", "Streamlit", "FAISS", "OpenAI", "Docling", "BeautifulSoup"]
x_start = Inches(3.0)
for i, tech in enumerate(techs):
    add_shape(slide, x_start + Inches(i * 1.25), Inches(5.3), Inches(1.1), Inches(0.45), BG_CARD)
    add_text(slide, x_start + Inches(i * 1.25), Inches(5.32), Inches(1.1), Inches(0.45),
             tech, 11, ACCENT, bold=True, alignment=PP_ALIGN.CENTER)


# ── Save ──
output_path = "/home/user/CDP_Chatbot/CDP_Chatbot_Pipeline_Comparison.pptx"
prs.save(output_path)
print(f"PPT saved to: {output_path}")
