"""
export_chat_review.py
=====================
Standalone, one-off script to export CDPChatHistory rows into an Excel file
for manual review.

- Reads ALL columns from CDPChatHistory between START_DATE and END_DATE
  (filtered on the CreatedOn column).
- Writes one .xlsx file with every column, plus one extra blank column at the
  end ("Review") so a reviewer can mark each row Correct / Incorrect.

This script is intentionally kept SEPARATE from the main service: it carries
its own connection string and its own configuration block below. It does not
import settings.py or anything from CDPAssist.py.

Usage:
    python export_chat_review.py

Requires: pyodbc, openpyxl
    pip install pyodbc openpyxl
"""

import sys
import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================
# CONFIGURATION  — edit these values, then run the script
# ============================================================

# --- Date range (inclusive). Format: YYYY-MM-DD ---------------
START_DATE = "2026-01-01"
END_DATE   = "2026-05-31"

# --- CDP2000 database connection (kept here on purpose,
#     independent of the main service settings) ---------------
DB_DRIVER   = "ODBC Driver 17 for SQL Server"
DB_SERVER   = "YOUR_CDP2000_SERVER"
DB_NAME     = "YOUR_CDP2000_DATABASE"
DB_USER     = "YOUR_USERNAME"
DB_PASSWORD = "YOUR_PASSWORD"

# --- Output ---------------------------------------------------
OUTPUT_FILE = "CDPChatHistory_review.xlsx"

# Table and the date column the range is filtered on
TABLE_NAME  = "CDPChatHistory"
DATE_COLUMN = "CreatedOn"

# Label for the extra blank review column added at the end
REVIEW_COLUMN_HEADER = "Review (Correct / Incorrect)"

# ============================================================
# Script — usually no need to edit below this line
# ============================================================


def get_connection():
    """Open a connection to the CDP2000 database (read-only usage)."""
    return pyodbc.connect(
        "DRIVER=" + DB_DRIVER + ";"
        "SERVER=" + DB_SERVER + ";"
        "DATABASE=" + DB_NAME + ";"
        "UID=" + DB_USER + ";"
        "PWD=" + DB_PASSWORD + ";"
    )


def fetch_rows():
    """
    Fetch all columns from the table within the date range.
    END_DATE is treated as inclusive of the whole day (< END_DATE + 1 day),
    so rows created at any time on END_DATE are included.
    """
    conn = get_connection()
    try:
        cursor = conn.cursor()
        query = (
            f"SELECT * FROM {TABLE_NAME} "
            f"WHERE {DATE_COLUMN} >= ? "
            f"AND {DATE_COLUMN} < DATEADD(DAY, 1, CAST(? AS DATE)) "
            f"ORDER BY {DATE_COLUMN} ASC"
        )
        cursor.execute(query, (START_DATE, END_DATE))
        columns = [col[0] for col in cursor.description]
        rows = cursor.fetchall()
        return columns, rows
    finally:
        conn.close()


def build_excel(columns, rows):
    """Write the rows to an .xlsx file with a trailing blank Review column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chat Review"

    headers = list(columns) + [REVIEW_COLUMN_HEADER]

    # Header row styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Data rows
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            # pyodbc returns native python types; datetimes need to be
            # tz-naive for openpyxl, which SQL Server DATETIME already is.
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Dropdown (Correct / Incorrect) on the review column for every data row
    review_col_idx = len(headers)
    review_col_letter = ws.cell(row=1, column=review_col_idx).column_letter
    if rows:
        dv = DataValidation(
            type="list",
            formula1='"Correct,Incorrect"',
            allow_blank=True,
        )
        dv.add(f"{review_col_letter}2:{review_col_letter}{len(rows) + 1}")
        ws.add_data_validation(dv)

    # Reasonable column widths
    for col_idx, header in enumerate(headers, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        name = header.lower()
        if name in ("question", "answer", "suggestedanswer"):
            ws.column_dimensions[letter].width = 60
        elif name == REVIEW_COLUMN_HEADER.lower():
            ws.column_dimensions[letter].width = 22
        else:
            ws.column_dimensions[letter].width = 20

    # Freeze the header row
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)


def main():
    print(f"Exporting {TABLE_NAME} from {START_DATE} to {END_DATE} (inclusive)...")
    try:
        columns, rows = fetch_rows()
    except Exception as e:
        print(f"ERROR: failed to read from the database: {e}")
        sys.exit(1)

    try:
        build_excel(columns, rows)
    except Exception as e:
        print(f"ERROR: failed to write the Excel file: {e}")
        sys.exit(1)

    print(f"Done. {len(rows)} row(s) written to '{OUTPUT_FILE}'.")
    print(f"Columns exported: {', '.join(columns)} + '{REVIEW_COLUMN_HEADER}'")


if __name__ == "__main__":
    main()
